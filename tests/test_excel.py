from openpyxl import load_workbook

from controle_receitas.excel import ABA_DETALHE, ABA_RESUMO, CABECALHO, gerar_planilha
from controle_receitas.models import Recebimento


def _exemplos():
    return [
        Recebimento(5, "Ana", 100.0, "1x2", 200.0),
        Recebimento(5, "Bruno", 50.0, "1x1", 50.0, pago="OK"),
        Recebimento(10, "Carla", 80.0, "2x3", 160.0),
    ]


def test_gera_planilha_com_cabecalho_e_dados(tmp_path):
    destino = gerar_planilha(_exemplos(), tmp_path / "saida" / "r.xlsx")
    ws = load_workbook(destino)[ABA_DETALHE]

    assert [c.value for c in ws[1]] == CABECALHO
    assert ws.max_row == 4
    assert ws["C2"].number_format == "R$ #,##0.00"


def test_resumo_agrupa_por_dia(tmp_path):
    destino = gerar_planilha(_exemplos(), tmp_path / "r.xlsx")
    resumo = load_workbook(destino)[ABA_RESUMO]

    # dia, nº, parcelas do mês, já recebido, falta receber, total do parcelado
    assert [c.value for c in resumo[2]] == [5, 2, 150.0, 50.0, 100.0, 250.0]
    assert [c.value for c in resumo[3]] == [10, 1, 80.0, 0.0, 80.0, 160.0]


def test_adiciona_em_planilha_existente(tmp_path):
    destino = tmp_path / "r.xlsx"
    gerar_planilha(_exemplos(), destino)
    gerar_planilha(_exemplos()[:1], destino)

    ws = load_workbook(destino)[ABA_DETALHE]
    assert ws.max_row == 5  # cabeçalho + 3 + 1
