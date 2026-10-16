"""Geração da planilha Excel formatada."""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import Recebimento, marcado_como_pago

ABA_DETALHE = "Recebimentos"
ABA_RESUMO = "Resumo por dia"

CABECALHO = ["Dia", "Cliente", "Valor da parcela", "Parcelas", "Total a receber", "Observação", "Pago"]
COLUNAS_DINHEIRO = ("C", "E")
FORMATO_DINHEIRO = "R$ #,##0.00"

_ESTILO_CABECALHO = {
    "font": Font(bold=True, color="FFFFFF"),
    "fill": PatternFill("solid", fgColor="1F3A93"),
    "alignment": Alignment(horizontal="center", vertical="center"),
}


def _estilizar_cabecalho(ws: Worksheet, titulos: list[str]) -> None:
    ws.append(titulos)
    for celula in ws[1]:
        celula.font = _ESTILO_CABECALHO["font"]
        celula.fill = _ESTILO_CABECALHO["fill"]
        celula.alignment = _ESTILO_CABECALHO["alignment"]
    ws.freeze_panes = "A2"


def _ajustar_larguras(ws: Worksheet) -> None:
    for coluna in ws.columns:
        maior = max(len(str(c.value)) if c.value is not None else 0 for c in coluna)
        ws.column_dimensions[get_column_letter(coluna[0].column)].width = min(maior + 4, 40)


def _abrir_ou_criar(caminho: Path) -> Workbook:
    """Abre a planilha existente (preservando dados) ou cria uma nova com cabeçalho."""
    if caminho.exists() and caminho.stat().st_size > 0:
        wb = load_workbook(caminho)
        if ABA_DETALHE in wb.sheetnames:
            return wb
        ws = wb.create_sheet(ABA_DETALHE, 0)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = ABA_DETALHE
    _estilizar_cabecalho(ws, CABECALHO)
    return wb


def _montar_resumo(wb: Workbook) -> None:
    """Reconstrói a aba de resumo por dia.

    Colunas: nº de recebimentos, valor das parcelas que entram no mês, quanto
    já foi recebido (coluna "Pago" marcada), quanto falta receber e o total
    do parcelado (todas as parcelas do intervalo).
    """
    if ABA_RESUMO in wb.sheetnames:
        del wb[ABA_RESUMO]
    resumo = wb.create_sheet(ABA_RESUMO)
    _estilizar_cabecalho(
        resumo,
        [
            "Dia",
            "Nº de recebimentos",
            "Parcelas do mês",
            "Já recebido",
            "Falta receber",
            "Total do parcelado",
        ],
    )

    detalhe = wb[ABA_DETALHE]
    # dia -> [quantidade, parcelas do mês, já recebido, total parcelado]
    por_dia: dict[int, list[float]] = defaultdict(lambda: [0, 0.0, 0.0, 0.0])
    for dia, _, valor_parcela, _, total, _, pago in detalhe.iter_rows(
        min_row=2, values_only=True
    ):
        if dia is None:
            continue
        valor = float(valor_parcela or 0)
        acumulado = por_dia[int(dia)]
        acumulado[0] += 1
        acumulado[1] += valor
        acumulado[2] += valor if marcado_como_pago(pago) else 0.0
        acumulado[3] += float(total or 0)

    for dia in sorted(por_dia):
        quantidade, parcelas_mes, recebido, total_parcelado = por_dia[dia]
        resumo.append(
            [
                dia,
                quantidade,
                round(parcelas_mes, 2),
                round(recebido, 2),
                round(parcelas_mes - recebido, 2),
                round(total_parcelado, 2),
            ]
        )

    ultima = resumo.max_row
    resumo.append(
        ["Total"] + [f"=SUM({letra}2:{letra}{ultima})" for letra in "BCDEF"]
    )
    for celula in resumo[resumo.max_row]:
        celula.font = Font(bold=True)
    for letra in "CDEF":
        for celula in resumo[letra][1:]:
            celula.number_format = FORMATO_DINHEIRO
    _ajustar_larguras(resumo)


def gerar_planilha(recebimentos: list[Recebimento], caminho: Path) -> Path:
    """Adiciona os recebimentos à planilha (criando-a se necessário) e a salva."""
    caminho = Path(caminho)
    caminho.parent.mkdir(parents=True, exist_ok=True)

    wb = _abrir_ou_criar(caminho)
    ws = wb[ABA_DETALHE]

    for recebimento in recebimentos:
        ws.append(recebimento.como_linha())

    for celula in ws["A"][1:]:
        celula.number_format = "0"
    for letra in COLUNAS_DINHEIRO:
        for celula in ws[letra][1:]:
            celula.number_format = FORMATO_DINHEIRO
    _ajustar_larguras(ws)

    _montar_resumo(wb)
    wb.save(caminho)
    return caminho
